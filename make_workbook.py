"""
Build Pricing-October-Data.xlsx: a clean, formatted Excel view of the database.

The CSVs in data/ are machine-written logs; opening and saving them in Excel
rewrites every row and breaks the pipeline (it happened on 2026-10-07). This
workbook is the thing to open instead. It is regenerated from data/pricing.db,
so run build_db.py first.

Run:  python build_db.py && python make_workbook.py
"""

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent
DB = ROOT / "data" / "pricing.db"
OUT = ROOT / "Pricing-October-Data.xlsx"

FONT = "Arial"
NAVY = "0B1F3A"

# (sheet name, table name, SQL, {column: number format}, [column widths])
SHEETS = [
    ("Results", "Results", """
        SELECT series AS "Series", CAST(game_no AS INT) AS "Game", date AS "Date", home AS "Home", away AS "Away",
               CASE WHEN played = 1 THEN 'yes' ELSE 'no' END AS "Played",
               CAST(home_w_before AS INT) || '-' || CAST(home_l_before AS INT) AS "Home record before",
               winner AS "Winner", note AS "Note"
        FROM results ORDER BY date, series, game_no""",
        {}, [9, 6, 11, 22, 22, 7, 17, 8, 60]),

    ("Latest Prices", "LatestPrices", """
        SELECT title AS "Game", substr(start_utc, 1, 16) AS "First pitch (UTC)",
               substr(captured_at, 1, 16) AS "Captured (UTC)", ROUND(hours_to_game, 1) AS "Hours to game",
               CAST(listing_count AS INT) AS "Listings", CAST(ticket_count AS INT) AS "Tickets",
               get_in_price AS "Get-in $", median_price AS "Median $", avg_price AS "Average $", max_price AS "Max $"
        FROM resale r
        WHERE captured_at = (SELECT MAX(captured_at) FROM resale r2 WHERE r2.event_id = r.event_id)
          AND bucket = 'mlb_postseason'
        ORDER BY start_utc""",
        {"Get-in $": "$#,##0", "Median $": "$#,##0", "Average $": "$#,##0", "Max $": "$#,##0", "Hours to game": "0.0",
         "Listings": "#,##0", "Tickets": "#,##0"},
        [62, 18, 18, 13, 10, 10, 10, 10, 11, 10]),

    ("Captures", "Captures", """
        SELECT title AS "Game", substr(start_utc, 1, 16) AS "First pitch (UTC)",
               substr(captured_at, 1, 16) AS "Captured (UTC)", ROUND(hours_to_game, 1) AS "Hours to game",
               CAST(listing_count AS INT) AS "Listings", CAST(ticket_count AS INT) AS "Tickets",
               get_in_price AS "Get-in $", median_price AS "Median $", avg_price AS "Average $", max_price AS "Max $",
               bucket AS "Bucket"
        FROM resale ORDER BY start_utc, title, captured_at""",
        {"Get-in $": "$#,##0", "Median $": "$#,##0", "Average $": "$#,##0", "Max $": "$#,##0", "Hours to game": "0.0",
         "Listings": "#,##0", "Tickets": "#,##0"},
        [62, 18, 18, 13, 10, 10, 10, 10, 11, 10, 15]),

    ("Games", "Games", """
        SELECT title AS "Game", substr(datetime_utc, 1, 16) AS "First pitch (UTC)", substr(datetime_local, 1, 16) AS "Local time",
               CASE WHEN time_tbd = 'True' THEN 'yes' ELSE '' END AS "Time TBD",
               home_team AS "Home", away_team AS "Away", venue AS "Venue", city AS "City",
               substr(visible_until_utc, 1, 16) AS "Listed until (UTC)"
        FROM seatgeek
        WHERE captured_at = (SELECT MAX(captured_at) FROM seatgeek) AND bucket = 'mlb_postseason'
        ORDER BY datetime_utc""",
        {}, [62, 18, 18, 9, 22, 22, 26, 14, 18]),

    ("Ticketmaster Status", "TicketmasterStatus", """
        SELECT name AS "Event", local_date AS "Date", local_time AS "Time", status AS "Status", venue AS "Venue", city AS "City",
               substr(sale_start, 1, 16) AS "Public sale start (UTC)", substr(sale_end, 1, 16) AS "Sale end (UTC)"
        FROM ticketmaster
        WHERE captured_at = (SELECT MAX(captured_at) FROM ticketmaster) AND bucket = 'mlb_postseason'
        ORDER BY local_date, name""",
        {}, [60, 11, 9, 11, 26, 14, 22, 18]),
]

README = [
    ("Pricing October — data workbook", True),
    ("A formatted view of data/pricing.db. Regenerate any time with:  python build_db.py && python make_workbook.py", False),
    ("", False),
    ("Do not open or save the CSV files in data/ with Excel. They are machine-written logs; saving them from Excel", False),
    ("rewrites every row and breaks the hourly pipeline. Open this workbook instead.", False),
    ("", False),
    ("Sheets", True),
    ("Results — one row per game, hand-kept from the bracket. 'Home record before' is the home team's series record going into the game.", False),
    ("Latest Prices — the most recent capture of each postseason game: listings, get-in, median, average, max.", False),
    ("Captures — every capture of every game (one row per game per capture). This is the table the SQL queries use.", False),
    ("Games — every postseason game SeatGeek currently lists, with UTC first pitch and the time it stops being listed.", False),
    ("Ticketmaster Status — the clubs' own listings and whether each is onsale, offsale or cancelled.", False),
    ("", False),
    ("Not included: the 779,000+ individual listings (section, row, price). They live in data/listings.csv and in the", False),
    ("'listings' table of the database; query them in SQL rather than in Excel.", False),
    ("", False),
    ("Grain of Captures: one row per game per capture. Captures happen at about 72, 48, 24, 12, 6, 3 and 1 hour before first pitch.", False),
    ("Prices are resale asks (what sellers list), not sales. Median is the number to quote; average is skewed by premium seats.", False),
]


def write_sheet(wb, name, table_name, rows, headers, formats, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    n_rows, n_cols = len(rows) + 1, len(headers)
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
        ws.cell(row=1, column=c).font = Font(name=FONT, bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).alignment = Alignment(vertical="center")
        fmt = formats.get(headers[c - 1])
        for rr in range(2, n_rows + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name=FONT, size=10)
            if fmt:
                cell.number_format = fmt
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "B2"
    ref = f"A1:{get_column_letter(n_cols)}{max(n_rows, 2)}"
    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)
    return ws


def main():
    conn = sqlite3.connect(DB)
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 130
    for i, (text, bold) in enumerate(README, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11, color=NAVY if bold else "000000")
    for name, table_name, sql, formats, widths in SHEETS:
        cur = conn.execute(sql)
        headers = [d[0] for d in cur.description]
        rows = cur.fetchall()
        write_sheet(wb, name, table_name, rows, headers, formats, widths)
        print(f"{name}: {len(rows)} rows")
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
